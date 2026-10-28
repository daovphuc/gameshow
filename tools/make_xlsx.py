"""Build sheet-template/gameshow-template.xlsx from the sheet-template CSVs. Usage: python -X utf8 tools/make_xlsx.py ."""
import csv, sys
from pathlib import Path
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill

repo = Path(sys.argv[1])
src = repo / "sheet-template"
out = src / "gameshow-template.xlsx"

COLS = {
    "teams": {
        "id": "Mã đội, duy nhất, không dấu (vd: doi-01). Bắt buộc.",
        "name": "Tên đội hiển thị. Bắt buộc.",
        "image_url": "Tên file ảnh đội trong thư mục images/ (vd: doi-01.jpg). Trống thì dùng chữ viết tắt tên đội.",
        "color": "Màu avatar, mã hex (tùy chọn), vd #D97757.",
    },
    "round1": {
        "package": "Tên gói chủ đề (mỗi gói = 1 câu hỏi/chủ đề). Bắt buộc.",
        "team_id": "Mã đội sở hữu gói (vd: doi-01). Bắt buộc.",
        "question": "Nội dung chủ đề tự giới thiệu / thuyết trình. Bắt buộc.",
        "time_seconds": "Thời gian tối đa (giây). Trống = dùng round1_time_seconds trong config.",
        "points": "Điểm tối đa của phần thi. Trống = dùng round1_default_points.",
        "image_url": "Tên file ảnh minh họa trong thư mục images/ (tùy chọn).",
    },
    "round2": {
        "set": "Mã bộ đề (vd: 1, 2...). Bắt buộc.",
        "number": "Số thứ tự câu hỏi trong bộ đề (1-10). Bắt buộc.",
        "question": "Nội dung câu hỏi trắc nghiệm. Bắt buộc.",
        "option_a": "Nội dung đáp án A. Bắt buộc.",
        "option_b": "Nội dung đáp án B. Bắt buộc.",
        "option_c": "Nội dung đáp án C. Bắt buộc.",
        "option_d": "Nội dung đáp án D. Bắt buộc.",
        "correct": "Đáp án đúng (A, B, C hoặc D). Bắt buộc.",
        "time_seconds": "Thời gian trả lời (giây). Trống = dùng round2_time_seconds trong config.",
        "image_url": "Tên file ảnh minh họa trong thư mục images/ (tùy chọn).",
    },
    "round3": {
        "number": "Số ô tình huống trên bảng (1-15). Bắt buộc, không trùng.",
        "question": "Nội dung tình huống nghiệp vụ PCCC và CNCH. Bắt buộc.",
        "time_seconds": "Thời gian trả lời (giây). Trống = dùng round3_time_seconds trong config.",
        "points": "Điểm tối đa của ô. Trống = dùng round3_default_points trong config.",
        "image_url": "Tên file ảnh minh họa trong thư mục images/ (tùy chọn).",
    },
    "judges": {
        "id": "Mã giám khảo, duy nhất (vd: gk-01). Bắt buộc.",
        "name": "Họ tên hoặc danh xưng giám khảo. Bắt buộc.",
    },
    "config": {
        "key": "Tên cài đặt (key). Không đổi tên cột key.",
        "value": "Giá trị cài đặt (tên hội thi, thời gian, điểm tối đa, màu giao diện...).",
    },
    "accounts": {
        "username": "Tên đăng nhập (không phân biệt hoa thường). Bắt buộc, duy nhất.",
        "password": "Mật khẩu đăng nhập (văn bản thuần). Bắt buộc.",
        "role": "Vai trò: host (dẫn chương trình), judge (giám khảo), hoặc team (đội thi). Bắt buộc.",
        "ref_id": "Mã giám khảo (từ tab judges) hoặc mã đội (từ tab teams). Để trống với host.",
        "display_name": "Tên hiển thị người dùng (tùy chọn).",
    },
}

HEAD_FILL = PatternFill("solid", fgColor="D97757")
HEAD_FONT = Font(name="Arial", bold=True, color="FFFFFF")
BODY_FONT = Font(name="Arial")

wb = Workbook()
guide = wb.active
guide.title = "huong_dan"
lines = [
    "HƯỚNG DẪN SỬ DỤNG FILE MẪU HỘI THI PCCC VÀ CNCH",
    "",
    "1. Vào Google Drive > Mới > Tải tệp lên > chọn file này. Mở file > Tệp > Lưu dưới dạng Google Trang tính.",
    "2. Giữ nguyên TÊN CÁC TAB (teams, round1, round2, round3, judges, accounts, config) và TÊN CỘT ở dòng 1 (ô màu cam).",
    "3. Sửa dữ liệu từ dòng 2 trở đi. Rê chuột vào tiêu đề cột để xem ghi chú hướng dẫn từng cột.",
    "4. Chia sẻ sheet: Bất kỳ ai có đường liên kết > Người xem. Copy link, dán vào mục Cài đặt trên màn hình Dẫn chương trình.",
    "5. Điểm số và trạng thái trận đấu được đồng bộ thời gian thực qua Firebase Realtime Database.",
    "6. Sửa sheet xong, bấm 'Tải lại dữ liệu' trên web để cập nhật toàn bộ thiết bị.",
    "",
    "Dữ liệu trong các tab là ví dụ mẫu cho 12 đội, 3 giám khảo, 2 bộ đề trắc nghiệm Vòng 2 và 15 tình huống Vòng 3.",
]
for i, text in enumerate(lines, 1):
    c = guide.cell(row=i, column=1, value=text)
    c.font = Font(name="Arial", bold=(i == 1), size=14 if i == 1 else 11)
guide.column_dimensions["A"].width = 110

for tab in ["teams", "round1", "round2", "round3", "judges", "accounts", "config"]:
    path = src / f"{tab}.csv"
    with open(path, encoding="utf-8-sig", newline="") as f:
        rows = [r for r in csv.reader(f) if any(cell.strip() for cell in r)]
    ws = wb.create_sheet(tab)
    header = rows[0]
    widths = [len(h) + 4 for h in header]
    for r_i, row in enumerate(rows, 1):
        for c_i, val in enumerate(row, 1):
            if r_i > 1 and val.strip() != "" and header[c_i - 1] in ("order", "time_seconds", "points", "number"):
                try:
                    val = int(val) if float(val).is_integer() else float(val)
                except ValueError:
                    pass
            cell = ws.cell(row=r_i, column=c_i, value=val)
            if r_i == 1:
                cell.font, cell.fill = HEAD_FONT, HEAD_FILL
                note = COLS.get(tab, {}).get(val)
                if note:
                    cell.comment = Comment(note, "Mẫu")
            else:
                cell.font = BODY_FONT
                cell.alignment = Alignment(vertical="top", wrap_text=len(str(val)) > 60)
            widths[c_i - 1] = max(widths[c_i - 1], min(len(str(val)) + 2, 70))
    for c_i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=c_i).column_letter].width = w
    ws.freeze_panes = "A2"
    print(tab, len(rows) - 1, "rows", header)

wb.save(out)
print("saved", out)
